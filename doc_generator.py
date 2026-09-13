# -*- coding: utf-8 -*-
"""Word文档生成模块"""
import os
import random
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# 字号对应磅值
FONT_SIZE_POINTS = {
    "初号": 42, "小初": 36, "一号": 26, "小一": 24,
    "二号": 22, "小二": 18, "三号": 16, "小三": 15,
    "四号": 14, "小四": 12, "五号": 10.5, "小五": 9,
    "六号": 7.5, "小六": 6.5, "七号": 5.5, "八号": 5
}

BORDER_THICKNESS_MAP = {
    "0.5磅": 4, "1.0磅": 8, "1.5磅": 12,
    "2.0磅": 16, "2.5磅": 20, "3.0磅": 24
}

BORDER_STYLE_MAP = {
    "实线": "single", "虚线": "dashed", "点线": "dotted",
    "双线": "double", "无框": "nil"
}


class DocGenerator:
    """Word文档生成器"""

    def __init__(self):
        # 默认格式选项
        self.row_height = "1.5"
        self.english_font = "Arial"
        self.chinese_font = "微软雅黑"
        self.font_size = "四号"
        self.border_thickness = "1.0磅"
        self.border_style = "实线"
        self.font_weight = "常规"

    def get_font_size_points(self):
        """获取字号对应的磅值"""
        return FONT_SIZE_POINTS.get(self.font_size, 12)

    def _apply_font_to_run(self, run, font_type='auto'):
        """为run应用字体设置"""
        if font_type == 'english':
            run.font.name = self.english_font
        elif font_type == 'chinese':
            run.font.name = self.chinese_font
            run._element.rPr.rFonts.set(qn('w:eastAsia'), self.chinese_font)
        else:  # auto
            run.font.name = self.english_font
            run._element.rPr.rFonts.set(qn('w:eastAsia'), self.chinese_font)
        run.font.size = Pt(self.get_font_size_points())
        run.bold = (self.font_weight == "加粗")

    def _set_table_borders(self, table):
        """设置表格边框"""
        tbl = table._tbl
        tblPr = tbl.tblPr
        tblBorders = OxmlElement('w:tblBorders')
        sz = str(BORDER_THICKNESS_MAP.get(self.border_thickness, 8))
        val = BORDER_STYLE_MAP.get(self.border_style, "single")
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), val)
            border.set(qn('w:sz'), sz)
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), '000000')
            tblBorders.append(border)
        tblPr.append(tblBorders)

    def _apply_row_height(self, row, height_cm):
        """设置行高"""
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        trHeight = OxmlElement('w:trHeight')
        trHeight.set(qn('w:val'), str(int(height_cm * 567)))  # 1cm = 567 twips
        trHeight.set(qn('w:hRule'), 'exact')
        trPr.append(trHeight)

    def _add_date_header(self, doc, date_text):
        """添加日期标题"""
        if not date_text:
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(date_text)
        run.font.name = self.chinese_font
        run.font.size = Pt(22)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), self.chinese_font)
        run.bold = (self.font_weight == "加粗")

    def _set_cell_text(self, cell, text, font_type='auto'):
        """设置单元格文本和字体"""
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = para.add_run(text)
        self._apply_font_to_run(run, font_type)

    def generate_word_files(self, words, folder_path, ordered=True, date_str=None):
        """生成单词默写文档（释义默写版+单词默写版+答案版）"""
        if not words:
            return None

        selected_words = words.copy()
        if not ordered:
            random.shuffle(selected_words)

        suffix = "顺序" if ordered else "乱序"
        row_height_cm = float(self.row_height)

        # 处理日期
        date_for_filename = ""
        date_for_doc = ""
        if date_str:
            try:
                parts = date_str.split("-")
                if len(parts) == 3:
                    y, m, d = parts
                    date_for_filename = f"{int(y)}_{int(m)}_{int(d)}_"
                    date_for_doc = f"{int(y)}-{int(m)}-{int(d)}"
            except Exception:
                pass

        saved_files = []

        # ===== 释义默写版（英文在左，释义空） =====
        doc1 = Document()
        self._add_date_header(doc1, date_for_doc)
        h1 = doc1.add_heading('释义默写版', 0)
        for run in h1.runs:
            self._apply_font_to_run(run, 'chinese')

        table1 = doc1.add_table(rows=1, cols=4)
        table1.style = 'Table Grid'
        for col in table1.columns:
            col.width = Inches(1.5)
        self._set_table_borders(table1)

        row_idx = 0
        row_cells = None
        for i, word in enumerate(selected_words):
            if i % 2 == 0:
                row_cells = table1.add_row().cells
                row_idx += 1
            cell_idx = (i % 2) * 2
            self._set_cell_text(row_cells[cell_idx], word['word'], 'english')
            self._set_cell_text(row_cells[cell_idx + 1], "", 'auto')
            self._apply_row_height(table1.rows[row_idx], row_height_cm)

            if (i + 1) % 20 == 0 and i + 1 < len(selected_words):
                doc1.add_page_break()
                table1 = doc1.add_table(rows=1, cols=4)
                table1.style = 'Table Grid'
                for col in table1.columns:
                    col.width = Inches(1.5)
                self._set_table_borders(table1)
                row_idx = 0

        f1 = os.path.join(folder_path, f'{date_for_filename}释义默写版_{suffix}.docx')
        doc1.save(f1)
        saved_files.append(f1)

        # ===== 单词默写版（释义在左，英文空） =====
        doc2 = Document()
        self._add_date_header(doc2, date_for_doc)
        h2 = doc2.add_heading('单词默写版', 0)
        for run in h2.runs:
            self._apply_font_to_run(run, 'chinese')

        table2 = doc2.add_table(rows=1, cols=4)
        table2.style = 'Table Grid'
        for col in table2.columns:
            col.width = Inches(1.5)
        self._set_table_borders(table2)

        row_idx = 0
        row_cells = None
        for i, word in enumerate(selected_words):
            if i % 2 == 0:
                row_cells = table2.add_row().cells
                row_idx += 1
            cell_idx = (i % 2) * 2
            self._set_cell_text(row_cells[cell_idx], word['meaning'], 'chinese')
            self._set_cell_text(row_cells[cell_idx + 1], "", 'auto')
            self._apply_row_height(table2.rows[row_idx], row_height_cm)

            if (i + 1) % 20 == 0 and i + 1 < len(selected_words):
                doc2.add_page_break()
                table2 = doc2.add_table(rows=1, cols=4)
                table2.style = 'Table Grid'
                for col in table2.columns:
                    col.width = Inches(1.5)
                self._set_table_borders(table2)
                row_idx = 0

        f2 = os.path.join(folder_path, f'{date_for_filename}单词默写版_{suffix}.docx')
        doc2.save(f2)
        saved_files.append(f2)

        # ===== 答案版 =====
        doc3 = Document()
        self._add_date_header(doc3, date_for_doc)
        h3 = doc3.add_heading('答案版', 0)
        for run in h3.runs:
            self._apply_font_to_run(run, 'chinese')

        table3 = doc3.add_table(rows=1, cols=4)
        table3.style = 'Table Grid'
        for col in table3.columns:
            col.width = Inches(1.5)
        self._set_table_borders(table3)

        row_idx = 0
        row_cells = None
        for i, word in enumerate(selected_words):
            if i % 2 == 0:
                row_cells = table3.add_row().cells
                row_idx += 1
            cell_idx = (i % 2) * 2
            self._set_cell_text(row_cells[cell_idx], word['word'], 'english')
            self._set_cell_text(row_cells[cell_idx + 1], word['meaning'], 'chinese')
            self._apply_row_height(table3.rows[row_idx], row_height_cm)

            if (i + 1) % 20 == 0 and i + 1 < len(selected_words):
                doc3.add_page_break()
                table3 = doc3.add_table(rows=1, cols=4)
                table3.style = 'Table Grid'
                for col in table3.columns:
                    col.width = Inches(1.5)
                self._set_table_borders(table3)
                row_idx = 0

        f3 = os.path.join(folder_path, f'{date_for_filename}答案_{suffix}.docx')
        doc3.save(f3)
        saved_files.append(f3)

        return saved_files
